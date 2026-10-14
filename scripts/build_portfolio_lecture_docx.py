#!/usr/bin/env python3
"""Generate Portfolio Management lecture Word document."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_document_styles(doc):
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    return h


def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        run = p.add_run(bold_prefix)
        run.bold = True
        p.add_run(text)
    else:
        p.add_run(text)
    return p


def add_numbered(doc, text):
    return doc.add_paragraph(text, style="List Number")


def add_quote(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.right_indent = Inches(0.25)
    run = p.add_run(text)
    run.italic = True
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)


def add_page_break(doc):
    doc.add_page_break()


def add_code_block(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.35)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(10)


def title_page(doc, title, subtitle):
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t_run = title_p.add_run(title)
    t_run.bold = True
    t_run.font.size = Pt(22)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub.add_run(subtitle)
    sub_run.font.size = Pt(12)
    sub_run.italic = True
    add_page_break(doc)


def build():
    doc = Document()
    set_document_styles(doc)

    # Title page
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t_run = title.add_run(
        "Portfolio Management Lecture:\n"
        "From the “20-Something Modelur” Trap to Robust Strategy Ecosystems"
    )
    t_run.bold = True
    t_run.font.size = Pt(22)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub.add_run(
        "A brain-dump of theories and frameworks for capital allocation,\n"
        "cross-strategy diversification, and antifragile portfolio design"
    )
    sub_run.font.size = Pt(12)
    sub_run.italic = True

    doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("Lecture notes · Quantitative trading & portfolio construction\n")
    meta.add_run("Research synthesis with references")

    add_page_break(doc)

    # Executive summary
    add_heading(doc, "Executive Summary", 1)
    doc.add_paragraph(
        "Many early-career traders and quantitative researchers fall into what this lecture calls "
        "the “20-something Modelur” trap: over-modeling, over-optimizing, and over-fitting a single "
        "strategy or asset until the backtest looks pristine—while remaining fully exposed to "
        "external return curves, regime shifts, and estimation error. The antidote is not more "
        "complexity on one equity curve, but a portfolio of strategies with distinct drivers, "
        "horizons, and risk exposures."
    )
    doc.add_paragraph(
        "This document synthesizes classical portfolio theory, modern quantitative intuition, and "
        "practical heuristics for building systems that survive, adapt, and compound. The goal is "
        "long-term stability through diversification of process and exposure—not perfection of "
        "a single fragile bet."
    )

    add_page_break(doc)

    # Section 1
    add_heading(doc, "1. The “20-Something Modelur” Mindset", 1)
    add_heading(doc, "1.1 What the trap looks like", 2)
    doc.add_paragraph(
        "The Modelur trap is the belief that sharper code, more parameters, and tighter in-sample "
        "fit will translate into durable live profits. In practice, each additional degree of freedom "
        "increases the chance that you are describing noise rather than structure."
    )
    add_bullet(doc, " — Chasing the highest in-sample Sharpe by tuning thresholds, lookbacks, and filters.", bold_prefix="Over-fitting one regime")
    add_bullet(doc, " — Assuming stationarity because the last three years “proved” the idea.", bold_prefix="Single-asset obsession")
    add_bullet(doc, " — Treating correlation during calm markets as permanent diversification.", bold_prefix="Ignoring tail dependence")
    add_bullet(doc, " — Confusing statistical significance with economic edge after many trials.", bold_prefix="Multiple testing blindness")

    add_heading(doc, "1.2 Why elegant models still fail", 2)
    doc.add_paragraph(
        "Even a theoretically sound strategy remains a claim about future joint distributions of "
        "returns, volatilities, and correlations. You do not control macro shocks, liquidity events, "
        "policy shifts, or the behavior of other participants. Your model is one small map of a "
        "moving territory."
    )
    add_quote(
        doc,
        "Michaud (1989) famously described mean-variance optimization as “error maximization”: "
        "the optimizer overweights assets (or signals) with the largest estimated returns—precisely "
        "where estimation error is largest.",
    )

    add_heading(doc, "1.3 The portfolio-level escape", 2)
    doc.add_paragraph(
        "Instead of perfecting one fragile curve, design multiple strategies that fail for different "
        "reasons at different times. Think in terms of a dynamic ecosystem: momentum sleeves, "
        "mean-reversion sleeves, carry, volatility, macro overlays, and execution alphas—each with "
        "explicit risk budgets and kill criteria."
    )
    add_bullet(doc, " — Does it have an economic story, not only a statistical fit?", bold_prefix="Economic mechanism")
    add_bullet(doc, " — Does performance survive costs, slippage, and capacity constraints?", bold_prefix="Implementation realism")
    add_bullet(doc, " — Does it add independent risk, not a repackaged version of what you already run?", bold_prefix="Orthogonality")
    add_bullet(doc, " — What would make you stop trading it before it destroys capital?", bold_prefix="Pre-mortem")

    add_page_break(doc)

    # Section 2
    add_heading(doc, "2. The Rule of 30 and the Illusion of Diversification", 1)
    add_heading(doc, "2.1 Where “30” comes from", 2)
    doc.add_paragraph(
        "Classic mean-variance intuition suggests that idiosyncratic risk falls quickly as you add "
        "uncorrelated holdings. Under simplified assumptions, roughly 30 independent assets can "
        "remove most diversifiable variance. Arena Investors (2022) notes that in textbook "
        "setups, diversifiable variance may be reduced by about 97% near 30 positions—hence the "
        "folk “Rule of 30.”"
    )
    doc.add_paragraph(
        "Nomura research (cited in industry summaries) similarly finds that spread between best "
        "and worst portfolios narrows rapidly until roughly 20–25 stocks, with marginal benefit "
        "fading past about 40 holdings for many equity mandates."
    )

    add_heading(doc, "2.2 Why 30 is not enough in the real world", 2)
    doc.add_paragraph(
        "Investopedia’s “illusion of diversification” argument is that counting names is not the same "
        "as diversifying risk. Surz & Price show that even 60 stocks may capture only ~86% of "
        "market diversification (R² vs. a broad benchmark)—because true diversification requires "
        "exposure across sectors, styles, geographies, and asset classes, not merely more tickers "
        "in one region."
    )
    add_bullet(doc, " — Returns are fat-tailed; correlations spike in stress.", bold_prefix="Non-normality")
    add_bullet(doc, " — Concentrated winners dominate outcomes (FOMO risk).", bold_prefix="Selection risk")
    add_bullet(doc, " — Strategies share hidden factors (e.g., dollar liquidity, rates beta).", bold_prefix="Factor overlap")

    add_heading(doc, "2.3 Diminishing returns—and hard ceilings", 2)
    doc.add_paragraph(
        "Recent work on signal portfolios (Hentschel, 2025) emphasizes that with positive average "
        "correlation among signals, Sharpe improvements from “breadth” converge to a finite limit. "
        "Adding correlated variants does not keep boosting risk-adjusted returns indefinitely; it "
        "increases overlap. Diversification helps until it doesn’t—then you need genuinely new "
        "drivers or lower average correlation."
    )

    add_heading(doc, "2.4 Practical takeaway for quants", 2)
    doc.add_paragraph(
        "Use the Rule of 30 as a starting heuristic for idiosyncratic risk within one sleeve—not "
        "as proof that 30 correlated crypto alt strategies equal diversification. For systematic "
        "traders, the analog is: 30 parameter sets on the same idea ≠ 30 ideas."
    )

    add_page_break(doc)

    # Section 3 MPT
    add_heading(doc, "3. Modern Portfolio Theory (MPT) — What Still Matters", 1)
    doc.add_paragraph(
        "Markowitz’s Modern Portfolio Theory (Corporate Finance Institute summary) frames investors "
        "as risk-averse: for a given expected return, prefer lower variance. MPT’s enduring insight "
        "is that portfolio risk depends on covariances, not just individual volatilities."
    )
    add_heading(doc, "3.1 Idiosyncratic vs. systematic risk", 2)
    add_bullet(doc, " — Asset-specific; reducible via diversification.", bold_prefix="Idiosyncratic")
    add_bullet(doc, " — Market-wide; not diversifiable away.", bold_prefix="Systematic")

    add_heading(doc, "3.2 Efficient frontier and capital allocation line", 2)
    doc.add_paragraph(
        "The efficient frontier is the set of portfolios maximizing expected return for each level of "
        "risk. Introducing a risk-free asset yields the Capital Allocation Line (CAL); the tangency "
        "portfolio maximizes the Sharpe ratio (reward per unit of volatility)."
    )
    doc.add_paragraph(
        "MPT is a language for trade-offs. Its weakness is not the geometry—it is that inputs "
        "(means, variances, correlations) are estimated with error and regimes change."
    )

    add_page_break(doc)

    # Section 4 Sharpe
    add_heading(doc, "4. Optimizing for Sharpe Ratio — Power and Pitfalls", 1)
    add_heading(doc, "4.1 Definition and intuition", 2)
    doc.add_paragraph(
        "Sharpe ratio = (portfolio return − risk-free rate) / portfolio volatility. It summarizes "
        "risk-adjusted performance and appears as the slope of the CAL in mean-variance diagrams."
    )

    add_heading(doc, "4.2 Where most people go wrong", 2)
    add_bullet(
        doc,
        " In-sample max Sharpe portfolios are optimistically biased; expected out-of-sample Sharpe "
        "is lower, especially when assets ≫ observations (De Miguel et al.; Kan & Zhou).",
        bold_prefix="In-sample worship",
    )
    add_bullet(
        doc,
        " Small errors in means get amplified; Michaud’s “error maximization” overweight tails of "
        "the return estimate distribution.",
        bold_prefix="Mean estimation dominance",
    )
    add_bullet(
        doc,
        " Research (e.g., Kolm et al., 2022) shows correlation misspecification can dominate "
        "mean errors in distorting the perceived efficient frontier.",
        bold_prefix="Correlation instability",
    )
    add_bullet(
        doc,
        " Optimized weights that look optimal on history often underperform simple 1/N or "
        "risk-parity baselines out of sample.",
        bold_prefix="Complexity without regularization",
    )
    add_bullet(
        doc,
        " Selecting the best of many backtests inflates reported Sharpe; use Deflated Sharpe, "
        "PBO/CSCV, and walk-forward validation (Bailey & López de Prado).",
        bold_prefix="Multiple testing",
    )

    add_heading(doc, "4.3 Better practices", 2)
    add_numbered(doc, "Separate research Sharpe from investable Sharpe (net of fees, turnover, capacity).")
    add_numbered(doc, "Prefer stable weights: equal risk contribution, volatility targeting, or constrained optimization.")
    add_numbered(doc, "Shrink covariances (Ledoit–Wolf) and treat expected returns as noisy—or omit them.")
    add_numbered(doc, "Report distributions of outcomes under resampled correlations and stressed scenarios.")
    add_numbered(doc, "Optimize the portfolio of strategies, not only parameters within one strategy.")

    add_page_break(doc)

    # Section 5 Factor models
    add_heading(doc, "5. Factor Models and Cross-Strategy Diversification", 1)
    doc.add_paragraph(
        "Factor thinking asks: what systematic exposures explain returns? For a quant portfolio, "
        "each strategy is a bundle of factor bets—momentum, value, carry, volatility risk premium, "
        "liquidity provision, etc."
    )
    add_heading(doc, "5.1 Why factors beat counting strategies", 2)
    doc.add_paragraph(
        "Two strategies with different names may both lose when rates rise or when cross-asset "
        "correlation goes to one. Factor decomposition reveals hidden overlap before live drawdowns do."
    )
    add_heading(doc, "5.2 Building orthogonal sleeves", 2)
    add_bullet(doc, " — Align P&L on a common calendar and currency.", bold_prefix="Align returns")
    add_bullet(doc, " — Measure pairwise correlation and rolling correlation in stress windows.", bold_prefix="Correlation matrix")
    add_bullet(doc, " — Track simultaneous drawdowns, not just average correlation.", bold_prefix="Tail co-movement")
    add_bullet(doc, " — Require economic distinction + low correlation (< ~0.3) before adding a sleeve.", bold_prefix="Admission rule")

    add_heading(doc, "5.3 Simple combination rules that work", 2)
    doc.add_paragraph(
        "Empirical studies on multi-signal equity models often show equal-weight composites beating "
        "optimized weights out of sample: optimized weights fit historical correlations; equal weights "
        "are harder to overfit. The same logic applies to combining trading strategies."
    )

    add_page_break(doc)

    # Section 6 Robust portfolios
    add_heading(doc, "6. Portfolios That Survive When Code or Assumptions Fail", 1)
    add_heading(doc, "6.1 Design for failure modes", 2)
    doc.add_paragraph(
        "Assume any single model will be wrong eventually. Robust portfolios encode redundancy: "
        "multiple uncorrelated edges, position limits, volatility scaling, and circuit breakers."
    )
    add_bullet(doc, " — Cap weight per strategy and per factor cluster.", bold_prefix="Risk budgeting")
    add_bullet(doc, " — Scale exposure when realized vol rises.", bold_prefix="Vol targeting")
    add_bullet(doc, " — Predefine max loss per sleeve per month/quarter.", bold_prefix="Drawdown governance")
    add_bullet(doc, " — Paper or reduce size when live vs. backtest diverges.", bold_prefix="Monitoring")

    add_heading(doc, "6.2 Hierarchical Risk Parity (HRP)", 2)
    doc.add_paragraph(
        "Marcos López de Prado’s Hierarchical Risk Parity (2016) addresses “Markowitz’s curse”: as "
        "correlations rise, you need diversification more—but covariance inversion becomes unstable "
        "(high condition number), so mean-variance solutions become fragile."
    )
    add_numbered(doc, "Hierarchical clustering of assets/strategies by correlation distance.")
    add_numbered(doc, "Quasi-diagonalization of the covariance matrix (reorder without full PCA change of basis).")
    add_numbered(doc, "Recursive bisection: allocate capital inversely to cluster variance.")

    doc.add_paragraph(
        "HRP uses covariance structure without requiring fragile return forecasts and remains usable "
        "when the covariance matrix is ill-conditioned—common in large strategy portfolios."
    )

    add_heading(doc, "6.3 Antifragility in practice", 2)
    doc.add_paragraph(
        "Antifragile here means benefiting from disorder modestly: volatility sleeves that earn in "
        "chaos, rebalancing that buys dislocation, and research pipelines that learn from live "
        "errors. The portfolio should get more conservative when estimates are unreliable—not more levered."
    )

    add_page_break(doc)

    # Section 7 Heuristics
    add_heading(doc, "7. Simple Heuristics for a Smoother Return Curve", 1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Heuristic"
    hdr[1].text = "Intent"
    rows = [
        ("1/N or equal risk weight across strategies", "Reduce estimation error; baseline hard to beat OOS"),
        ("Cap pairwise correlation for new sleeves", "Avoid duplicate bets"),
        ("Volatility-scale each strategy before blending", "Prevent one loud sleeve from dominating"),
        ("Rebalance on bands (e.g., ±10% from target)", "Harvest diversification without over-trading"),
        ("Stress-test: 2008, 2020, rate shock scenarios", "See correlation breakdown before live pain"),
        ("Walk-forward + holdout years", "Separate research from validation"),
        ("Deflated Sharpe / PBO when trialing many variants", "Correct for selection bias"),
        ("Document kill rules in writing", "Remove discretion during drawdowns"),
    ]
    for h, intent in rows:
        row = table.add_row().cells
        row[0].text = h
        row[1].text = intent

    doc.add_paragraph()
    add_heading(doc, "7.1 Volatility vs. fragility", 2)
    doc.add_paragraph(
        "Lowering volatility alone does not create robustness if tail risks are synchronized. Prefer "
        "metrics that punish simultaneous losses: maximum portfolio drawdown under stress, "
        "downside correlation, and time-under-water at the portfolio level."
    )

    add_page_break(doc)

    # Section 8 Workflow
    add_heading(doc, "8. Suggested Workflow: From Single Strategy to Ecosystem", 1)
    add_numbered(doc, "Articulate economic mechanism and failure scenarios for Strategy A.")
    add_numbered(doc, "Validate with costs, capacity, and out-of-sample windows—not only IS Sharpe.")
    add_numbered(doc, "Add Strategy B only if drivers differ and stress correlations are acceptable.")
    add_numbered(doc, "Combine with simple weights first; introduce optimization only with shrinkage/constraints.")
    add_numbered(doc, "Monitor factor overlap monthly; retire sleeves that drift into redundancy.")
    add_numbered(doc, "Revisit capital allocation when regime indicators shift (vol level, correlation structure).")

    add_heading(doc, "8.1 Lecture closing", 2)
    add_quote(
        doc,
        "The aim isn’t to chase perfection—it’s to design systems that survive, adapt, and compound. "
        "See your portfolio as a dynamic ecosystem, not a single bet.",
    )

    add_page_break(doc)

    # Appendix A — Math
    add_heading(doc, "Appendix A: Core Formulas", 1)
    doc.add_paragraph(
        "These relationships underpin the lecture. They are intentionally compact; full derivations "
        "appear in standard texts (Markowitz; Cochrane; López de Prado)."
    )
    add_heading(doc, "A.1 Portfolio return and variance (two assets)", 2)
    add_code_block(
        doc,
        "E[R_p] = w_A · E[R_A] + w_B · E[R_B]\n\n"
        "Var(R_p) = w_A² σ_A² + w_B² σ_B² + 2 w_A w_B ρ_AB σ_A σ_B\n\n"
        "When ρ_AB < 1, portfolio variance is below the weighted average of individual variances.",
    )
    add_heading(doc, "A.2 Sharpe ratio", 2)
    add_code_block(
        doc,
        "SR = (E[R_p] − R_f) / σ_p\n\n"
        "On the mean-variance diagram, the tangency portfolio maximizes SR among risky portfolios.",
    )
    add_heading(doc, "A.3 Diversification and correlation", 2)
    doc.add_paragraph(
        "For N equally weighted assets with identical variance σ² and pairwise correlation ρ, "
        "portfolio variance scales approximately as:"
    )
    add_code_block(doc, "σ_p² ≈ (ρ + (1 − ρ) / N) · σ²")
    doc.add_paragraph(
        "As N → ∞ with fixed ρ > 0, variance does not go to zero—systematic correlation ρ remains. "
        "This is why “more names” without lower ρ hits a ceiling."
    )

    add_page_break(doc)

    # Appendix B — Worked example
    add_heading(doc, "Appendix B: Worked Example — Three Strategy Portfolio", 1)
    doc.add_paragraph(
        "Illustrative monthly statistics (education only—not live trade recommendations). "
        "Three sleeves: trend (T), mean reversion (M), volatility carry (V)."
    )
    table2 = doc.add_table(rows=1, cols=5)
    table2.style = "Table Grid"
    h2 = table2.rows[0].cells
    h2[0].text = "Sleeve"
    h2[1].text = "Ann. return"
    h2[2].text = "Ann. vol"
    h2[3].text = "Sharpe (≈)"
    h2[4].text = "Max DD (illustrative)"
    for row_data in [
        ("Trend (T)", "12%", "18%", "0.67", "−22%"),
        ("Mean rev. (M)", "8%", "12%", "0.67", "−15%"),
        ("Vol carry (V)", "6%", "10%", "0.60", "−12%"),
    ]:
        r = table2.add_row().cells
        for i, val in enumerate(row_data):
            r[i].text = val

    doc.add_paragraph()
    doc.add_paragraph("Pairwise correlations (monthly, illustrative):")
    table3 = doc.add_table(rows=1, cols=4)
    table3.style = "Table Grid"
    h3 = table3.rows[0].cells
    h3[0].text = ""
    h3[1].text = "T"
    h3[2].text = "M"
    h3[3].text = "V"
    for row_data in [("T", "1.00", "−0.25", "0.15"), ("M", "−0.25", "1.00", "0.10"), ("V", "0.15", "0.10", "1.00")]:
        r = table3.add_row().cells
        for i, val in enumerate(row_data):
            r[i].text = val

    doc.add_paragraph()
    add_heading(doc, "B.1 Equal weight vs. optimized weight", 2)
    doc.add_paragraph(
        "Equal weight (33/33/33): portfolio vol often falls below the average sleeve vol because "
        "T and M offset each other. A mean-variance optimizer fit on the same history might push "
        "80% into the sleeve with the highest estimated return—raising in-sample Sharpe but often "
        "hurting the next year when correlations shift."
    )
    add_heading(doc, "B.2 Stress overlay", 2)
    doc.add_paragraph(
        "Suppose in a crisis month, ρ_TM rises from −0.25 to +0.55 and all sleeves lose together. "
        "Lesson: monitor rolling stress correlation, not only full-sample averages. A portfolio "
        "that looks diversified in calm data can be one bet in a crash."
    )

    add_page_break(doc)

    # Appendix C — Agenda
    add_heading(doc, "Appendix C: 90-Minute Lecture Agenda", 1)
    agenda = [
        ("0:00–0:10", "Hook: best backtest vs. first live month; define Modelur trap"),
        ("0:10–0:25", "MPT refresher: idiosyncratic vs. systematic; efficient frontier"),
        ("0:25–0:40", "Rule of 30 & illusion of diversification; finite breadth limits"),
        ("0:40–0:55", "Sharpe optimization pitfalls; estimation error & multiple testing"),
        ("0:55–1:10", "Factor view + cross-strategy correlation; stress co-drawdowns"),
        ("1:10–1:20", "HRP and robust heuristics; risk budgeting"),
        ("1:20–1:30", "Q&A; homework: factor map + one add/cut decision"),
    ]
    t_agenda = doc.add_table(rows=1, cols=2)
    t_agenda.style = "Table Grid"
    t_agenda.rows[0].cells[0].text = "Time"
    t_agenda.rows[0].cells[1].text = "Segment"
    for time_seg, seg in agenda:
        row = t_agenda.add_row().cells
        row[0].text = time_seg
        row[1].text = seg

    add_page_break(doc)

    # Appendix D — Exercises
    add_heading(doc, "Appendix D: Discussion Questions & Exercises", 1)
    add_heading(doc, "Discussion questions", 2)
    add_numbered(doc, "When is a higher in-sample Sharpe evidence of skill vs. evidence of overfitting?")
    add_numbered(doc, "Why can two strategies with negative average correlation still blow up together?")
    add_numbered(doc, "What would convince you to remove a sleeve from the portfolio—metrics, not ego?")
    add_numbered(doc, "How does HRP differ from classic risk parity in how it treats correlation clusters?")

    add_heading(doc, "Take-home exercise", 2)
    doc.add_paragraph(
        "1. List every strategy or signal you run. 2. Assign each a primary driver (momentum, value, "
        "carry, vol, liquidity, macro, etc.). 3. Estimate pairwise correlation of weekly returns "
        "(or use your backtest logs). 4. Identify the most redundant pair. 5. Propose one new "
        "sleeve that would lower average pairwise correlation without doubling complexity."
    )

    add_page_break(doc)

    # References
    add_heading(doc, "References and Further Reading", 1)
    refs = [
        "Investopedia — “The Diversification Myth: Why 30 Stocks Aren’t Enough for Your Portfolio” "
        "(illusion of diversification): https://www.investopedia.com/articles/stocks/11/illusion-of-diversification.asp",
        "Corporate Finance Institute — Modern Portfolio Theory (MPT) overview: "
        "https://corporatefinanceinstitute.com/resources/wealth-management/modern-portfolio-theory-mpt/",
        "Wikipedia — Hierarchical Risk Parity (López de Prado, 2016): "
        "https://en.wikipedia.org/wiki/Hierarchical_Risk_Parity",
        "López de Prado, M. — “Building Diversified Portfolios that Outperform Out of Sample” (HRP).",
        "De Miguel, V., Garlappi, L., & Uppal, R. (2009) — “Optimal Versus Naive Diversification.”",
        "Michaud, R. (1989) — “The Markowitz Optimization Enigma: Is ‘Optimized’ Optimal?”",
        "Bailey, D. & López de Prado, M. (2014) — Backtest overfitting, Deflated Sharpe, PBO.",
        "Kolm, P. et al. (2022) — Effects of errors in means, variances, and correlations on MV framework.",
        "Hentschel, L. (2025) — “The Limits of Diversification” (finite Sharpe ceilings with correlated signals).",
        "Fidelity — Guide to diversification (rebalancing and monitoring): "
        "https://www.fidelity.com/viewpoints/investing-ideas/guide-to-diversification",
        "Markowitz, H. (1952) — Portfolio selection (foundational MPT).",
    ]
    for r in refs:
        doc.add_paragraph(r, style="List Bullet")

    add_heading(doc, "Speaker Notes (Expanded)", 1)
    notes = [
        ("Opening", "Personal story: pristine backtest, rough first live month. Frame lecture as escaping single-curve obsession."),
        ("Modelur trap", "Whiteboard: parameters ↑ → degrees of freedom ↑ → selection bias ↑. Mention Deflated Sharpe if audience is quant-heavy."),
        ("Rule of 30", "Contrast textbook 30-stock result with Investopedia ‘illusion’—sectors, geographies, styles matter."),
        ("Sharpe", "Show Michaud quote; emphasize correlation error dominance (Kolm et al.). Demo: equal weight beats optimized OOS."),
        ("Factors", "Align P&L on one calendar; show stress-month correlation spike (Appendix B)."),
        ("HRP", "Intuition only unless PhD crowd: clustering → bisection avoids matrix inversion blow-ups."),
        ("Close", "Assign Appendix D homework due next session."),
    ]
    for title, body in notes:
        p = doc.add_paragraph()
        run = p.add_run(f"{title}: ")
        run.bold = True
        p.add_run(body)

    out = "/workspace/lecture-materials/Portfolio_Management_Lecture.docx"
    doc.save(out)
    return out


def build_slides():
    doc = Document()
    set_document_styles(doc)
    title_page(
        doc,
        "Portfolio Management Lecture\n(Slide Deck)",
        "Companion to full lecture notes · ~25 slides",
    )

    slides = [
        (
            "Slide 1 — Title",
            [
                "From the 20-Something Modelur Trap to Robust Strategy Ecosystems",
                "Portfolio thinking for quant traders",
            ],
        ),
        (
            "Slide 2 — The problem",
            [
                "One perfect backtest ≠ one durable business",
                "External return curves dominate your P&L",
                "Optimization without breadth = fragile elegance",
            ],
        ),
        (
            "Slide 3 — Modelur trap",
            [
                "Over-model · over-optimize · over-fit",
                "Single asset / single regime",
                "Mistake in-sample Sharpe for edge",
            ],
        ),
        (
            "Slide 4 — The shift",
            [
                "Portfolio of strategies, not one strategy",
                "Different drivers · horizons · risks",
                "Survive · adapt · compound",
            ],
        ),
        (
            "Slide 5 — MPT in one slide",
            [
                "Risk-averse investors prefer less variance for same return",
                "Diversification cuts idiosyncratic risk",
                "Systematic risk remains",
            ],
        ),
        (
            "Slide 6 — Correlation",
            [
                "ρ = +1 → no diversification benefit",
                "ρ < 0 → offsets can stabilize P&L",
                "Estimate ρ poorly → portfolio breaks",
            ],
        ),
        (
            "Slide 7 — Efficient frontier",
            [
                "Max return per unit of risk",
                "Tangency portfolio → max Sharpe",
                "Inputs are estimated → outputs fragile",
            ],
        ),
        (
            "Slide 8 — Rule of 30",
            [
                "~30 names can remove much idiosyncratic risk (textbook)",
                "Real markets: fat tails, stress ρ, factor overlap",
                "30 correlated strategies ≠ diversified",
            ],
        ),
        (
            "Slide 9 — Illusion of diversification",
            [
                "Counting tickers ≠ diversifying exposures",
                "Need sectors, styles, geographies, asset classes",
                "Breadth has diminishing returns—and ceilings",
            ],
        ),
        (
            "Slide 10 — Sharpe ratio",
            [
                "SR = (R_p − R_f) / σ_p",
                "Good summary metric, dangerous optimization target",
            ],
        ),
        (
            "Slide 11 — Sharpe pitfalls",
            [
                "In-sample max SR is biased high",
                "Mean errors get amplified (error maximization)",
                "Many trials → pick lucky noise (multiple testing)",
            ],
        ),
        (
            "Slide 12 — Fixes",
            [
                "Shrink covariances · constrain weights",
                "Equal risk / 1/N baselines",
                "Walk-forward · Deflated Sharpe · PBO",
            ],
        ),
        (
            "Slide 13 — Factor lens",
            [
                "Each strategy = bundle of factor bets",
                "Hidden overlap kills diversification",
                "Decompose before you allocate",
            ],
        ),
        (
            "Slide 14 — Cross-strategy work",
            [
                "Align returns on common timeline",
                "Track stress co-drawdowns",
                "Admission rule: mechanism + low ρ",
            ],
        ),
        (
            "Slide 15 — Combining signals",
            [
                "Equal weight often beats optimized OOS",
                "Complexity fits history, not future",
            ],
        ),
        (
            "Slide 16 — Robust design",
            [
                "Assume any model fails eventually",
                "Risk budgets · vol targeting · kill rules",
            ],
        ),
        (
            "Slide 17 — Markowitz’s curse",
            [
                "More correlation → need diversification more",
                "Also → unstable covariance inversion",
            ],
        ),
        (
            "Slide 18 — HRP (3 steps)",
            [
                "1. Hierarchical cluster",
                "2. Quasi-diagonalize",
                "3. Recursive bisection",
            ],
        ),
        (
            "Slide 19 — Heuristics",
            [
                "Cap weights · rebalance bands · stress tests",
                "Document kill rules in advance",
            ],
        ),
        (
            "Slide 20 — Workflow",
            [
                "Validate sleeve → check orthogonality → simple blend",
                "Monitor factors → retire redundancy",
            ],
        ),
        (
            "Slide 21 — Worked example",
            [
                "Trend + mean reversion + vol carry",
                "Calm ρ vs. crisis ρ",
            ],
        ),
        (
            "Slide 22 — Antifragile mindset",
            [
                "Ecosystem, not single bet",
                "Lower leverage when estimates are noisy",
            ],
        ),
        (
            "Slide 23 — Homework",
            [
                "Factor-map your book",
                "Cut one redundant sleeve or propose one orthogonal add",
            ],
        ),
        (
            "Slide 24 — References",
            [
                "Investopedia · CFI MPT · Wikipedia HRP",
                "López de Prado · De Miguel · Bailey",
            ],
        ),
        (
            "Slide 25 — Q&A",
            ["What would break your portfolio tomorrow?"],
        ),
    ]

    for heading, bullets in slides:
        add_heading(doc, heading, 2)
        for b in bullets:
            add_bullet(doc, b)
        doc.add_paragraph()

    out = "/workspace/lecture-materials/Portfolio_Management_Lecture_Slides.docx"
    doc.save(out)
    return out


def build_handout():
    doc = Document()
    set_document_styles(doc)
    title_page(
        doc,
        "Student Handout",
        "Portfolio Management Lecture — key concepts (2-page summary)",
    )

    add_heading(doc, "Key ideas", 1)
    add_bullet(doc, " — optimize a portfolio of edges, not one over-fit curve.", bold_prefix="Think ecosystems")
    add_bullet(doc, " — idiosyncratic risk falls with uncorrelated bets; systematic risk does not.", bold_prefix="MPT")
    add_bullet(doc, " — counting holdings ≠ diversifying exposures.", bold_prefix="Rule of 30")
    add_bullet(doc, " — great in-sample, often disappointing out-of-sample.", bold_prefix="Sharpe max")
    add_bullet(doc, " — cluster correlations, allocate without fragile inversion.", bold_prefix="HRP")

    add_heading(doc, "Checklist before adding a strategy", 1)
    for item in [
        "Economic story independent of backtest",
        "Net of fees / slippage / capacity",
        "Pairwise correlation < ~0.3 with core book (or justified)",
        "Stress-month behavior reviewed",
        "Written kill rule (drawdown or drift)",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "Formulas to remember", 1)
    add_code_block(doc, "Var(R_p) = w' Σ w   |   SR = (E[R_p] − R_f) / σ_p")
    add_code_block(doc, "σ_p² ≈ (ρ + (1−ρ)/N) σ²  (equal-weight, identical σ, constant ρ)")

    add_heading(doc, "Further reading", 1)
    doc.add_paragraph(
        "Investopedia illusion of diversification; CFI Modern Portfolio Theory; "
        "Wikipedia Hierarchical Risk Parity."
    )

    out = "/workspace/lecture-materials/Portfolio_Management_Student_Handout.docx"
    doc.save(out)
    return out


if __name__ == "__main__":
    paths = [build(), build_slides(), build_handout()]
    for p in paths:
        print(p)

